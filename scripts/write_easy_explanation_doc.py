#!/usr/bin/env python3
"""Generate the current student explanation from README.md (requires python-docx)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

root = Path(__file__).resolve().parents[1]
doc = Document()
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(11)
code = False
for line in (root / 'README.md').read_text().splitlines():
    if line.startswith('```'):
        code = not code
        continue
    if not line.strip():
        continue
    if code:
        paragraph = doc.add_paragraph(line)
        for run in paragraph.runs:
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
    elif line.startswith('#'):
        level = len(line) - len(line.lstrip('#'))
        doc.add_heading(line.lstrip('# ').replace('**', ''), min(level, 3))
    else:
        text = line.replace('**', '').replace('`', '')
        doc.add_paragraph(text[2:] if text.startswith('- ') else text,
                          style='List Bullet' if text.startswith('- ') else None)
doc.save(root / 'WorkSure_Easy_Explanation.docx')
print('Updated WorkSure_Easy_Explanation.docx from README.md')
